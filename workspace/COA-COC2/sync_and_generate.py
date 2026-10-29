"""
COA-COC 数据同步与导入 Excel 生成工具

用法（在项目根目录运行）：
    python workspace/COA-COC2/sync_and_generate.py
    python workspace/COA-COC2/sync_and_generate.py --skip-fetch   # 跳过拉取，用本地 CSV

流程：
    1. 从金蝶拉取 BD_UNIT / QM_InspectMethod / QM_InspectItem /
                   QM_QCScheme / QM_QualityStd，保存到 data/{表单}_{时间戳}.csv
    2. 解析 input/test plan Excel，与金蝶数据对比，找出缺失项
    3. 生成 5 个导入 Excel 到 output/（含计量单位，格式与 template/ 下金蝶模板一致）

导入顺序（依赖关系）：
    1. 计量单位_*.xlsx
    2. 检验方法_*.xlsx
    3. 质量标准_*.xlsx
    4. 检验项目_*.xlsx   （引用检验方法、质量标准 FNumber）
    5. 质检方案_*.xlsx   （引用检验项目 FNumber）
"""
from __future__ import annotations

import argparse
import sys

if sys.stdout.encoding and sys.stdout.encoding.lower() != "utf-8":
    sys.stdout.reconfigure(encoding="utf-8")
if sys.stderr.encoding and sys.stderr.encoding.lower() != "utf-8":
    sys.stderr.reconfigure(encoding="utf-8")
from datetime import date, datetime
from pathlib import Path
from typing import Any

import openpyxl
import pandas as pd

# ── 路径 ───────────────────────────────────────────────────────────────────────
WORKSPACE    = Path(__file__).resolve().parent          # workspace/COA-COC2
PROJECT_ROOT = WORKSPACE.parent.parent                  # 项目根目录
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from src.config import get_settings
from src.kingdee_client import KingdeeClient

def _resolve_workspace_dir(primary: str, legacy: str) -> Path:
    """优先使用新目录名；若只存在旧目录名，则回退兼容。"""
    primary_path = WORKSPACE / primary
    legacy_path = WORKSPACE / legacy
    if primary_path.exists():
        return primary_path
    if legacy_path.exists():
        return legacy_path
    return primary_path


def _dir_label(path: Path) -> str:
    return f"{path.name}/"


TEMPLATE_DIR = WORKSPACE / "template"
INPUT_DIR    = _resolve_workspace_dir("input", "import")
DATA_DIR     = WORKSPACE / "data"
OUTPUT_DIR   = _resolve_workspace_dir("output", "export")

ORG_ID   = "100"
ORG_NAME = "斯顿食品原料（苏州）有限公司"
AUDIT_USER_ID = "SHENXP"
AUDIT_USER_NAME = "Shen Ami"
# 计量单位导入：与业务约定固定值（换算率等仍沿用模板样例行）
BD_UNIT_GROUP_ID   = 21
BD_UNIT_GROUP_NAME = "QA"
BD_UNIT_ROUND_TYPE = "四舍五入"
BD_UNIT_PRECISION  = 0
TODAY    = date.today().isoformat()
NOW      = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
RUN_TS   = datetime.now().strftime("%Y%m%d_%H%M%S")
EXPIRE   = "9999-12-31"

FORM_IDS = [
    "BD_UNIT",
    "QM_InspectMethod",
    "QM_InspectItem",
    "QM_QCScheme",
    "QM_QualityStd",
]


def _resolve_local_csv(form_id: str) -> Path | None:
    """本地缓存：优先取最新的 ``{form_id}_*.csv``（按修改时间），否则回退 ``{form_id}.csv``。"""
    stamped = list(DATA_DIR.glob(f"{form_id}_*.csv"))
    if stamped:
        return max(stamped, key=lambda p: p.stat().st_mtime)
    legacy = DATA_DIR / f"{form_id}.csv"
    return legacy if legacy.exists() else None


# ══════════════════════════════════════════════════════════════════════════════
# 1. 从金蝶拉取数据
# ══════════════════════════════════════════════════════════════════════════════

def _main_entry(entries: list[dict]) -> dict | None:
    """取主表实体（含 FBillHead 的 entry），没有则取第一个。"""
    for e in entries:
        if "FBillHead" in e.get("entry_key", ""):
            return e
    return entries[0] if entries else None


def fetch_form_data(client: KingdeeClient, form_id: str) -> pd.DataFrame:
    """
    1. QueryBusinessInfo → 取主表 FieldKeys
    2. ExecuteBillQuery 分页拉全量数据
    """
    print(f"  [字段] {form_id} ...", end=" ", flush=True)
    field_info = client.get_form_field_keys(form_id)
    entry = _main_entry(field_info.get("entries", []))
    field_keys: list[str] = entry["field_keys"] if entry else field_info["field_keys"]
    print(f"{len(field_keys)} 个字段")

    field_keys_csv = ",".join(field_keys)
    all_rows: list[Any] = []
    start_row, limit = 0, 2000

    while True:
        rows = client.execute_bill_query({
            "FormId":      form_id,
            "FieldKeys":   field_keys_csv,
            "FilterString": "",
            "OrderString":  "",
            "TopRowCount":  0,
            "StartRow":     start_row,
            "Limit":        limit,
            "SubSystemId":  "",
        })
        if not rows:
            break
        all_rows.extend(rows)
        if len(rows) < limit:
            break
        start_row += limit

    print(f"  [数据] {form_id}: 共 {len(all_rows)} 行")
    if all_rows:
        return pd.DataFrame(all_rows, columns=field_keys)
    return pd.DataFrame(columns=field_keys)


def _clear_data_dir_csv_cache() -> None:
    """拉取前删除 data 目录下全部 CSV，避免堆积多批历史快照。"""
    if not DATA_DIR.is_dir():
        return
    n = 0
    for p in DATA_DIR.glob("*.csv"):
        try:
            p.unlink(missing_ok=True)
            n += 1
        except OSError:
            pass
    if n:
        print(f"  [清理] 已删除 {n} 个旧 CSV 文件")


def _clear_output_dir_xlsx_cache() -> None:
    """生成导入文件前删除 output 目录下全部 xlsx，避免堆积历史导入包。"""
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    n = 0
    for p in OUTPUT_DIR.glob("*.xlsx"):
        try:
            p.unlink(missing_ok=True)
            n += 1
        except OSError:
            pass
    if n:
        print(f"  [清理] 已删除 {n} 个旧 xlsx 文件")


def fetch_and_save(skip_fetch: bool = False) -> dict[str, pd.DataFrame]:
    """拉取所有业务对象数据，保存为 CSV；skip_fetch=True 时读本地缓存。"""
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    dfs: dict[str, pd.DataFrame] = {}

    if skip_fetch:
        print("跳过拉取，从本地 CSV 加载 ...")
        for fid in FORM_IDS:
            p = _resolve_local_csv(fid)
            if p is not None:
                dfs[fid] = pd.read_csv(p, dtype=str, keep_default_na=False)
                print(f"  {fid}: {len(dfs[fid])} 行（本地 {p.name}）")
            else:
                print(f"  {fid}: 本地 CSV 不存在，跳过")
                dfs[fid] = pd.DataFrame()
        return dfs

    _clear_data_dir_csv_cache()
    print("从金蝶拉取数据 ...")
    with KingdeeClient(get_settings()) as client:
        for fid in FORM_IDS:
            df = fetch_form_data(client, fid)
            csv_path = DATA_DIR / f"{fid}_{RUN_TS}.csv"
            df.to_csv(csv_path, index=False, encoding="utf-8-sig")
            print(f"  → 保存 {csv_path.name}")
            dfs[fid] = df
    return dfs


# ══════════════════════════════════════════════════════════════════════════════
# 2. 工具函数
# ══════════════════════════════════════════════════════════════════════════════

_NAME_PREF = ("FName#2052", "FName#1033", "FName#3076", "FName")


def _name_col(df: pd.DataFrame) -> str | None:
    """找最优先的 FName 列名。"""
    for col in _NAME_PREF:
        if col in df.columns:
            return col
    return next((c for c in df.columns if "FName" in c), None)


def _norm(v: Any) -> str:
    """去首尾空白；None / nan / 'None' 均返回空字符串。"""
    s = str(v).strip() if v is not None else ""
    return "" if s in ("nan", "None", "NaN") else s


def _build_name_to_fnumber(df: pd.DataFrame) -> dict[str, str]:
    """构建 {name: FNumber} 查找表（name 用 _norm 标准化）。"""
    nc = _name_col(df)
    if df.empty or nc is None or "FNumber" not in df.columns:
        return {}
    result: dict[str, str] = {}
    for _, row in df.iterrows():
        k = _norm(row[nc])
        v = _norm(row["FNumber"])
        if k:
            result[k] = v
    return result


def _next_fnumbers(df: pd.DataFrame, prefix: str, count: int) -> list[str]:
    """从 CSV 中找最大 FNumber，返回 count 个新序号（prefix + 6位数字）。"""
    max_num = 0
    for val in (df["FNumber"] if "FNumber" in df.columns else []):
        v = _norm(val)
        if v.upper().startswith(prefix.upper()):
            try:
                max_num = max(max_num, int(v[len(prefix):]))
            except ValueError:
                pass
    return [f"{prefix}{max_num + i:06d}" for i in range(1, count + 1)]


def _next_unit_codes(df: pd.DataFrame, count: int) -> list[str]:
    """从 BD_UNIT CSV 找最大 QA-XXX 编码，返回 count 个新编码（QA-XXX 格式）。"""
    max_num = 0
    for val in (df["FNumber"] if "FNumber" in df.columns else []):
        v = _norm(val).upper()
        if v.startswith("QA-"):
            try:
                max_num = max(max_num, int(v[3:]))
            except ValueError:
                pass
    return [f"QA-{max_num + i:03d}" for i in range(1, count + 1)]


def _set(row: list, idx: dict[str, int], key: str, val: Any) -> None:
    """安全地给 row[key] 赋值（key 不存在时静默跳过）。"""
    if key in idx:
        row[idx[key]] = val


def _set_bill_head_seq(row: list, idx: dict[str, int], seq: int) -> None:
    """给当前 sheet 的 FBillHead 列写入从 1 开始的导入序号。"""
    bill_head_key = next((key for key in idx if key.startswith("FBillHead(")), None)
    if bill_head_key is not None:
        row[idx[bill_head_key]] = str(seq)


def _set_audit_fields(row: list, idx: dict[str, int]) -> None:
    """统一写入创建/修改/审核人与日期字段。"""
    _set(row, idx, "FCreatorId", AUDIT_USER_ID)
    _set(row, idx, "FCreatorId#Name", AUDIT_USER_NAME)
    _set(row, idx, "FCreateDate", NOW)
    _set(row, idx, "FModifierId", AUDIT_USER_ID)
    _set(row, idx, "FModifierId#Name", AUDIT_USER_NAME)
    _set(row, idx, "FModifyDate", NOW)
    _set(row, idx, "FApproverId", AUDIT_USER_ID)
    _set(row, idx, "FApproverId#Name", AUDIT_USER_NAME)
    _set(row, idx, "FApproveDate", NOW)


# ══════════════════════════════════════════════════════════════════════════════
# 3. 解析 Test Plan 并计算差异
# ══════════════════════════════════════════════════════════════════════════════

_REQUIRED_COLS = {"material_code", "material_name", "test_item", "target", "unit", "method"}


def parse_test_plan() -> pd.DataFrame:
    """
    从输入目录读取 test plan Excel（取第一个 xlsx）。
    自动检测含必要列（material_code / test_item / target / unit / method）的 sheet。
    """
    candidates = sorted(INPUT_DIR.glob("*.xlsx"))
    if not candidates:
        raise FileNotFoundError(f"{_dir_label(INPUT_DIR)} 目录下未找到任何 xlsx 文件：{INPUT_DIR}")
    path = candidates[0]
    if len(candidates) > 1:
        print(f"  [警告] {_dir_label(INPUT_DIR)} 下有多个 xlsx，使用：{path.name}")

    xl = pd.ExcelFile(path)
    candidates_found: list[tuple[int, str]] = []  # (priority, sheet_name)
    for sheet in xl.sheet_names:
        try:
            head = pd.read_excel(xl, sheet_name=sheet, nrows=1, dtype=str)
            cols = set(head.columns)
            if not _REQUIRED_COLS.issubset(cols):
                continue
            # 优先级：sheet 名含 "test plan" > 无 Unnamed/重复列 > 其他
            unnamed_count = sum(1 for c in cols if str(c).startswith("Unnamed"))
            name_match = 1 if "test plan" in sheet.lower() else 0
            priority = (name_match * 1000) - unnamed_count
            candidates_found.append((priority, sheet))
        except Exception:
            continue

    if not candidates_found:
        raise ValueError(
            f"{path.name} 中未找到含必要列 {_REQUIRED_COLS} 的 sheet，"
            f"现有 sheets：{xl.sheet_names}"
        )
    matched_sheet = max(candidates_found, key=lambda x: x[0])[1]

    df = pd.read_excel(xl, sheet_name=matched_sheet, dtype=str)
    for col in df.columns:
        df[col] = df[col].fillna("").str.strip()
    print(f"Test plan ({path.name} / {matched_sheet}): {len(df)} 行，列：{df.columns.tolist()}")
    return df


def compute_missing(tp: pd.DataFrame, dfs: dict[str, pd.DataFrame]) -> dict[str, pd.DataFrame]:
    """
    返回各类型缺失（新增）项的 DataFrame：
        QM_InspectMethod  列：method
        QM_QualityStd     列：target
        QM_InspectItem    列：test_item, unit, target, method
        QM_QCScheme       列：material_code, material_name, test_item, target, unit, method
    """
    missing: dict[str, pd.DataFrame] = {}

    # ── 计量单位 ──────────────────────────────────────────────────────────────
    df_u = dfs.get("BD_UNIT", pd.DataFrame())
    nc   = _name_col(df_u)
    existing_units = {_norm(v) for v in df_u[nc]} if (nc and not df_u.empty) else set()
    new_units = [u for u in tp["unit"].map(_norm).unique() if u and u not in existing_units]
    missing["BD_UNIT"] = pd.DataFrame({"unit": new_units})
    print(f"  新计量单位:  {len(new_units)}")

    # ── 检验方法 ──────────────────────────────────────────────────────────────
    df_m = dfs.get("QM_InspectMethod", pd.DataFrame())
    nc   = _name_col(df_m)
    existing_methods = {_norm(v) for v in df_m[nc]} if (nc and not df_m.empty) else set()
    new_methods = [m for m in tp["method"].map(_norm).unique() if m and m not in existing_methods]
    missing["QM_InspectMethod"] = pd.DataFrame({"method": new_methods})
    print(f"  新检验方法:  {len(new_methods)}")

    # ── 质量标准 ──────────────────────────────────────────────────────────────
    df_s = dfs.get("QM_QualityStd", pd.DataFrame())
    nc   = _name_col(df_s)
    existing_stds = {_norm(v) for v in df_s[nc]} if (nc and not df_s.empty) else set()
    new_stds = [s for s in tp["target"].map(_norm).unique() if s and s not in existing_stds]
    missing["QM_QualityStd"] = pd.DataFrame({"target": new_stds})
    print(f"  新质量标准:  {len(new_stds)}")

    # ── 检验项目 ──────────────────────────────────────────────────────────────
    df_i = dfs.get("QM_InspectItem", pd.DataFrame())
    nc   = _name_col(df_i)
    existing_items = {_norm(v) for v in df_i[nc]} if (nc and not df_i.empty) else set()
    new_item_rows = (
        tp[~tp["test_item"].map(_norm).isin(existing_items)]
        .drop_duplicates(subset=["test_item"])
        [["test_item", "unit", "target", "method"]]
        .copy()
    )
    missing["QM_InspectItem"] = new_item_rows
    print(f"  新检验项目:  {len(new_item_rows)}")

    # ── 质检方案 ──────────────────────────────────────────────────────────────
    df_sc = dfs.get("QM_QCScheme", pd.DataFrame())
    existing_schemes: set[str] = set()
    if "FNumber" in df_sc.columns and not df_sc.empty:
        existing_schemes = {_norm(v) for v in df_sc["FNumber"]}
    new_scheme_rows = (
        tp[~tp["material_code"].map(_norm).isin(existing_schemes)]
        .drop_duplicates(subset=["material_code", "test_item"])
        .copy()
    )
    missing["QM_QCScheme"] = new_scheme_rows
    n_materials = new_scheme_rows["material_code"].nunique()
    print(f"  新质检方案:  {n_materials} 个物料，{len(new_scheme_rows)} 行明细")

    return missing


# ══════════════════════════════════════════════════════════════════════════════
# 4. 读取模板 & 写出导入 Excel
# ══════════════════════════════════════════════════════════════════════════════

_TEMPLATE_PREFIX: dict[str, str] = {
    "BD_UNIT": "计量单位",
    "QM_InspectMethod": "检验方法",
    "QM_InspectItem": "检验项目",
    "QM_QCScheme": "质检方案",
    "QM_QualityStd": "质量标准",
}


def _template_path(form_id: str) -> Path:
    """匹配 ``{中文名}.xlsx`` 或 ``{中文名}_*.xlsx``（按文件名排序取第一个）。"""
    prefix = _TEMPLATE_PREFIX[form_id]
    matches = sorted(TEMPLATE_DIR.glob(f"{prefix}*.xlsx"))
    if not matches:
        raise FileNotFoundError(f"未找到 {form_id} 模板：{TEMPLATE_DIR / (prefix + '*.xlsx')}")
    return matches[0]


def _read_template_headers(form_id: str) -> tuple[str, list, list]:
    """
    读模板主数据 sheet 的前两行。
    返回 (sheet_name, h0, h1)。
    QM_InspectItem 有两个 sheet，取不含"分组"的主数据 sheet。
    """
    wb = openpyxl.load_workbook(_template_path(form_id), data_only=True)
    sheet_name = next((s for s in wb.sheetnames if "分组" not in s), wb.sheetnames[0])
    ws = wb[sheet_name]
    rows = list(ws.iter_rows(values_only=True))
    return sheet_name, list(rows[0]), list(rows[1])


def _save_import_excel(filename: str, sheet_name: str, h0: list, h1: list, data_rows: list[list]) -> None:
    """写出单 sheet 导入 Excel，保存到输出目录。"""
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    out = OUTPUT_DIR / filename
    wb  = openpyxl.Workbook()
    ws  = wb.active
    ws.title = sheet_name
    ws.append(h0)
    ws.append(h1)
    for row in data_rows:
        ws.append(row)
    wb.save(out)
    print(f"  → {out.name}：{len(data_rows)} 行新记录")


def _save_inspect_item_excel(filename: str, h0: list, h1: list, data_rows: list[list]) -> None:
    """
    写出检验项目导入 Excel（双 sheet）：
      sheet1：数据分组#单据头(FBillHead)Group  —— 从模板原样复制
      sheet2：检验项目#单据头(FBillHead)        —— 模板表头 + 新数据行
    """
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    out = OUTPUT_DIR / filename

    # 读模板的分组 sheet（原样保留）
    wb_tpl = openpyxl.load_workbook(_template_path("QM_InspectItem"), data_only=True)
    group_sheet_name = next((s for s in wb_tpl.sheetnames if "分组" in s), None)
    item_sheet_name  = next((s for s in wb_tpl.sheetnames if "分组" not in s), wb_tpl.sheetnames[0])

    wb = openpyxl.Workbook()

    # sheet1：分组（原样复制）
    if group_sheet_name:
        ws_group = wb.active
        ws_group.title = group_sheet_name
        for row in wb_tpl[group_sheet_name].iter_rows(values_only=True):
            ws_group.append(list(row))

    # sheet2：检验项目（新数据）
    ws_items = wb.create_sheet(title=item_sheet_name)
    ws_items.append(h0)
    ws_items.append(h1)
    for row in data_rows:
        ws_items.append(row)

    wb.save(out)
    print(f"  → {out.name}：{len(data_rows)} 行新记录（含分组 sheet）")


# ── 各类型生成函数 ─────────────────────────────────────────────────────────────

def gen_unit(missing_df: pd.DataFrame, dfs: dict[str, pd.DataFrame]) -> dict[str, str]:
    lu_unit = _build_name_to_fnumber(dfs.get("BD_UNIT", pd.DataFrame()))
    if missing_df.empty:
        print("  计量单位：无新增，跳过")
        return lu_unit

    df_u = dfs.get("BD_UNIT", pd.DataFrame())
    sheet_name, h0, h1 = _read_template_headers("BD_UNIT")
    idx = {v: i for i, v in enumerate(h0)}

    wb_tpl = openpyxl.load_workbook(_template_path("BD_UNIT"), data_only=True)
    ws_tpl = wb_tpl[sheet_name]
    sample_row = next(ws_tpl.iter_rows(min_row=3, max_row=3, values_only=True), None)
    base = list(sample_row) if sample_row else []
    while len(base) < len(h0):
        base.append(None)
    base = base[: len(h0)]

    new_codes = _next_unit_codes(df_u, len(missing_df))
    rows: list[list[Any]] = []
    for seq, (row, code) in enumerate(zip(missing_df.itertuples(), new_codes), start=1):
        unit_name = _norm(row.unit)
        r = list(base)
        r[0] = None
        _set_bill_head_seq(r, idx, seq)
        _set(r, idx, "FNumber", code)
        _set(r, idx, "FName#1033", unit_name)
        _set(r, idx, "FName#2052", unit_name)
        _set(r, idx, "FName#3076", unit_name)
        _set(r, idx, "FUnitGroupId", BD_UNIT_GROUP_ID)
        _set(r, idx, "FUnitGroupId#Name", BD_UNIT_GROUP_NAME)
        _set(r, idx, "FPrecision", BD_UNIT_PRECISION)
        _set(r, idx, "FRoundType", BD_UNIT_ROUND_TYPE)
        _set(r, idx, "SubHeadEntity", None)
        rows.append(r)
        lu_unit[unit_name] = code

    _save_import_excel(f"计量单位_{RUN_TS}_导入.xlsx", sheet_name, h0, h1, rows)
    return lu_unit


def gen_inspect_method(missing_df: pd.DataFrame, dfs: dict[str, pd.DataFrame]) -> dict[str, str]:
    lu_method = _build_name_to_fnumber(dfs.get("QM_InspectMethod", pd.DataFrame()))
    if missing_df.empty:
        print("  检验方法：无新增，跳过")
        return lu_method
    sheet_name, h0, h1 = _read_template_headers("QM_InspectMethod")
    idx = {v: i for i, v in enumerate(h0)}
    new_fnumbers = _next_fnumbers(dfs.get("QM_InspectMethod", pd.DataFrame()), "JYFF", len(missing_df))
    rows = []
    for seq, (row, fnumber) in enumerate(zip(missing_df.itertuples(), new_fnumbers), start=1):
        m = _norm(row.method)
        r = [None] * len(h0)
        r[0] = None
        _set_bill_head_seq(r, idx, seq)
        _set(r, idx, "FNumber",            fnumber)
        _set(r, idx, "FCreateOrgId",       ORG_ID)
        _set(r, idx, "FCreateOrgId#Name",  ORG_NAME)
        _set(r, idx, "FUseOrgId",          ORG_ID)
        _set(r, idx, "FUseOrgId#Name",     ORG_NAME)
        _set(r, idx, "FName#1033",         m)
        _set(r, idx, "FName#2052",         m)
        _set(r, idx, "FName#3076",         m)
        _set_audit_fields(r, idx)
        rows.append(r)
        lu_method[m] = fnumber
    _save_import_excel(f"检验方法_{RUN_TS}_导入.xlsx", sheet_name, h0, h1, rows)
    return lu_method


def gen_quality_std(missing_df: pd.DataFrame, dfs: dict[str, pd.DataFrame]) -> dict[str, str]:
    lu_std = _build_name_to_fnumber(dfs.get("QM_QualityStd", pd.DataFrame()))
    if missing_df.empty:
        print("  质量标准：无新增，跳过")
        return lu_std
    sheet_name, h0, h1 = _read_template_headers("QM_QualityStd")
    idx = {v: i for i, v in enumerate(h0)}
    new_fnumbers = _next_fnumbers(dfs.get("QM_QualityStd", pd.DataFrame()), "ZLBZ", len(missing_df))
    rows = []
    for seq, (row, fnumber) in enumerate(zip(missing_df.itertuples(), new_fnumbers), start=1):
        t = _norm(row.target)
        r = [None] * len(h0)
        r[0] = None
        _set_bill_head_seq(r, idx, seq)
        _set(r, idx, "FNumber",            fnumber)
        _set(r, idx, "FCreateOrgId",       ORG_ID)
        _set(r, idx, "FCreateOrgId#Name",  ORG_NAME)
        _set(r, idx, "FUseOrgId",          ORG_ID)
        _set(r, idx, "FUseOrgId#Name",     ORG_NAME)
        _set(r, idx, "FName#1033",         t)
        _set(r, idx, "FName#2052",         t)
        _set(r, idx, "FName#3076",         t)
        _set_audit_fields(r, idx)
        rows.append(r)
        lu_std[t] = fnumber
    _save_import_excel(f"质量标准_{RUN_TS}_导入.xlsx", sheet_name, h0, h1, rows)
    return lu_std


def gen_inspect_item(
    missing_df: pd.DataFrame,
    dfs:        dict[str, pd.DataFrame],
    lu_unit:    dict[str, str],
    lu_std:     dict[str, str],
    lu_method:  dict[str, str],
) -> dict[str, str]:
    lu_item = _build_name_to_fnumber(dfs.get("QM_InspectItem", pd.DataFrame()))
    if missing_df.empty:
        print("  检验项目：无新增，跳过")
        return lu_item
    _, h0, h1 = _read_template_headers("QM_InspectItem")
    idx = {v: i for i, v in enumerate(h0)}
    new_fnumbers = _next_fnumbers(dfs.get("QM_InspectItem", pd.DataFrame()), "JYXM", len(missing_df))
    rows = []
    for seq, (row, fnumber) in enumerate(zip(missing_df.itertuples(), new_fnumbers), start=1):
        ti     = _norm(row.test_item)
        unit   = _norm(row.unit)
        target = _norm(row.target)
        method = _norm(row.method)

        r = [None] * len(h0)
        r[0] = None
        _set_bill_head_seq(r, idx, seq)
        _set(r, idx, "FNumber",              fnumber)
        _set(r, idx, "FCreateOrgId",        ORG_ID)
        _set(r, idx, "FCreateOrgId#Name",   ORG_NAME)
        _set(r, idx, "FUseOrgId",           ORG_ID)
        _set(r, idx, "FUseOrgId#Name",      ORG_NAME)
        _set(r, idx, "FName#1033",          ti)
        _set(r, idx, "FName#2052",          ti)
        _set(r, idx, "FName#3076",          ti)
        _set(r, idx, "FUnitId",             lu_unit.get(unit) or None)
        _set(r, idx, "FUnitId#Name",        unit or None)
        _set(r, idx, "FAnalysisMethod",     "其他分析")
        _set(r, idx, "FDefectlevel",        "重缺陷")
        _set(r, idx, "FDestructInspect",    "False")
        _set(r, idx, "FKeyInspect",         "True")
        _set(r, idx, "FQualityStdId",       lu_std.get(target) or None)
        _set(r, idx, "FQualityStdId#Name",  target or None)
        _set(r, idx, "FInspectMethodId",    lu_method.get(method) or None)
        _set(r, idx, "FInspectMethodId#Name", method or None)
        _set(r, idx, "FEnableQcValue",      "False")
        _set(r, idx, "FIsZeroNotValidate",  "False")
        _set_audit_fields(r, idx)
        rows.append(r)
        lu_item[ti] = fnumber
    _save_inspect_item_excel(f"检验项目_{RUN_TS}_导入.xlsx", h0, h1, rows)
    return lu_item


def gen_qcscheme(
    missing_df: pd.DataFrame,
    dfs:        dict[str, pd.DataFrame],
    lu_item:    dict[str, str],
    lu_std:     dict[str, str],
    lu_method:  dict[str, str],
    lu_unit:    dict[str, str],
) -> None:
    if missing_df.empty:
        print("  质检方案：无新增，跳过")
        return
    sheet_name, h0, h1 = _read_template_headers("QM_QCScheme")
    idx = {v: i for i, v in enumerate(h0)}
    rows = []
    scheme_seq = 0

    for material_code, group in missing_df.groupby("material_code", sort=False):
        material_code = _norm(material_code)
        material_name = _norm(group["material_name"].iloc[0])
        scheme_seq += 1

        for item_seq, item in enumerate(group.itertuples(), 1):
            ti     = _norm(item.test_item)
            target = _norm(item.target)
            unit   = _norm(item.unit)
            method = _norm(item.method)

            r = [None] * len(h0)

            # 方案头：仅首行填写
            if item_seq == 1:
                r[0] = None
                _set_bill_head_seq(r, idx, scheme_seq)
                _set(r, idx, "FCreateOrgId",       ORG_ID)
                _set(r, idx, "FCreateOrgId#Name",  ORG_NAME)
                _set(r, idx, "FUseOrgId",          ORG_ID)
                _set(r, idx, "FUseOrgId#Name",     ORG_NAME)
                _set(r, idx, "FNumber",            material_code)
                _set(r, idx, "FName#1033",         material_name)
                _set(r, idx, "FName#2052",         material_name)
                _set(r, idx, "FName#3076",         material_name)
                _set(r, idx, "FExpireDate",        EXPIRE)
                _set(r, idx, "FEffectDate",        TODAY)
                _set(r, idx, "FApplyScope",        "通用")
                _set_audit_fields(r, idx)

            # 检验项目明细（FEntity 子表，每行都填）
            _set(r, idx, "FInspectItemId",           lu_item.get(ti) or None)
            _set(r, idx, "FInspectItemId#Name",      ti)
            _set(r, idx, "FAnalysisMethod",          "其他分析")
            _set(r, idx, "FDefectlevel",             "重缺陷")
            _set(r, idx, "FKeyInspect",              "True")
            _set(r, idx, "FDestructInspect",         "False")
            _set(r, idx, "FEnableQcValue",           "False")
            _set(r, idx, "FIsZeroNotValidate",       "False")
            _set(r, idx, "FInspectMethodId",         lu_method.get(method) or None)
            _set(r, idx, "FInspectMethodId#Name",    method or None)
            _set(r, idx, "FQualityStdId",            lu_std.get(target) or None)
            _set(r, idx, "FQualityStdId#Name",       target or None)
            _set(r, idx, "FStatus",                  "合格")
            _set(r, idx, "FUnitId",                  lu_unit.get(unit) or None)
            _set(r, idx, "FUnitId#Name",             unit or None)
            _set(r, idx, "FCompareSymbol",           "=")
            rows.append(r)

    _save_import_excel(f"质检方案_{RUN_TS}_导入.xlsx", sheet_name, h0, h1, rows)


# ══════════════════════════════════════════════════════════════════════════════
# 主流程
# ══════════════════════════════════════════════════════════════════════════════

def main() -> None:
    parser = argparse.ArgumentParser(description="COA-COC 数据同步与导入 Excel 生成")
    parser.add_argument("--skip-fetch", action="store_true",
                        help="跳过金蝶拉取，使用已有 data/*.csv")
    args = parser.parse_args()

    print("=" * 60)
    print("步骤 1：拉取金蝶数据")
    print("=" * 60)
    dfs = fetch_and_save(skip_fetch=args.skip_fetch)

    print("\n" + "=" * 60)
    print("步骤 2：解析 Test Plan")
    print("=" * 60)
    tp = parse_test_plan()

    print("\n" + "=" * 60)
    print("步骤 3：计算差异")
    print("=" * 60)
    missing = compute_missing(tp, dfs)

    print("\n" + "=" * 60)
    print("步骤 4：生成导入 Excel")
    print("=" * 60)
    _clear_output_dir_xlsx_cache()

    # 依次生成：每步返回已有+新生成的 name→FNumber 查找表，供下游使用
    lu_unit   = gen_unit(missing["BD_UNIT"], dfs)
    lu_method = gen_inspect_method(missing["QM_InspectMethod"], dfs)
    lu_std    = gen_quality_std(missing["QM_QualityStd"], dfs)
    lu_item   = gen_inspect_item(missing["QM_InspectItem"], dfs, lu_unit, lu_std, lu_method)
    gen_qcscheme(missing["QM_QCScheme"], dfs, lu_item, lu_std, lu_method, lu_unit)

    print("\n完成！")
    print(f"  数据 CSV：{DATA_DIR}")
    print(f"  导入 Excel：{OUTPUT_DIR}")
    print("\n建议导入顺序：")
    print(f"  1. 计量单位_{RUN_TS}_导入.xlsx")
    print(f"  2. 检验方法_{RUN_TS}_导入.xlsx")
    print(f"  3. 质量标准_{RUN_TS}_导入.xlsx")
    print(f"  4. 检验项目_{RUN_TS}_导入.xlsx")
    print(f"  5. 质检方案_{RUN_TS}_导入.xlsx")


if __name__ == "__main__":
    main()
