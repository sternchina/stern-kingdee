from __future__ import annotations

import importlib.util
import unittest
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook


SCRIPT_PATH = Path(__file__).resolve().parents[1] / "sync_and_generate.py"


def _load_module():
    spec = importlib.util.spec_from_file_location("coa_coc2_sync", SCRIPT_PATH)
    module = importlib.util.module_from_spec(spec)
    assert spec.loader is not None
    spec.loader.exec_module(module)
    return module


def _row_map(xlsx_path: Path, sheet_name: str) -> dict[str, object]:
    wb = load_workbook(xlsx_path, data_only=True)
    ws = wb[sheet_name]
    headers = [cell for cell in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
    values = [cell for cell in next(ws.iter_rows(min_row=3, max_row=3, values_only=True))]
    return dict(zip(headers, values))


def _row_maps(xlsx_path: Path, sheet_name: str) -> list[dict[str, object]]:
    wb = load_workbook(xlsx_path, data_only=True)
    ws = wb[sheet_name]
    headers = [cell for cell in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
    rows: list[dict[str, object]] = []
    for values in ws.iter_rows(min_row=3, values_only=True):
        if all(value is None for value in values):
            continue
        rows.append(dict(zip(headers, values)))
    return rows


def _assert_audit_fields(row: dict[str, object], current_ts: str) -> None:
    assert row["FCreatorId"] == "SHENXP"
    assert row["FCreatorId#Name"] == "Shen Ami"
    assert row["FCreateDate"] == current_ts
    assert row["FModifierId"] == "SHENXP"
    assert row["FModifierId#Name"] == "Shen Ami"
    assert row["FModifyDate"] == current_ts
    assert row["FApproverId"] == "SHENXP"
    assert row["FApproverId#Name"] == "Shen Ami"
    assert row["FApproveDate"] == current_ts


class AuditFieldTests(unittest.TestCase):
    def test_generators_fill_unified_audit_fields(self):
        mod = _load_module()
        fixed_today = "2026-04-13"
        fixed_now = "2026-04-13 16:20:30"
        fixed_run_ts = "20990101_000000"
        output_dir = SCRIPT_PATH.parent / "output"
        created_files = [
            output_dir / f"计量单位_{fixed_run_ts}_导入.xlsx",
            output_dir / f"检验方法_{fixed_run_ts}_导入.xlsx",
            output_dir / f"质量标准_{fixed_run_ts}_导入.xlsx",
            output_dir / f"检验项目_{fixed_run_ts}_导入.xlsx",
            output_dir / f"质检方案_{fixed_run_ts}_导入.xlsx",
        ]

        original_output_dir = mod.OUTPUT_DIR
        original_today = mod.TODAY
        original_now = getattr(mod, "NOW", None)
        original_run_ts = mod.RUN_TS
        try:
            mod.OUTPUT_DIR = output_dir
            mod.TODAY = fixed_today
            mod.NOW = fixed_now
            mod.RUN_TS = fixed_run_ts

            for path in created_files:
                path.unlink(missing_ok=True)

            mod.gen_unit(
                pd.DataFrame([{"unit": "mg/kg"}, {"unit": "g/L"}]),
                {"BD_UNIT": pd.DataFrame([{"FNumber": "QA-001", "FName#2052": "Old Unit"}])},
            )
            unit_rows = _row_maps(
                output_dir / f"计量单位_{fixed_run_ts}_导入.xlsx",
                "计量单位#单据头(FBillHead)",
            )
            self.assertEqual(["1", "2"], [row["FBillHead(BD_UNIT)"] for row in unit_rows])

            mod.gen_inspect_method(
                pd.DataFrame([{"method": "Method A"}, {"method": "Method B"}]),
                {"QM_InspectMethod": pd.DataFrame([{"FNumber": "JYFF000100", "FName#2052": "Old Method"}])},
            )
            method_rows = _row_maps(
                output_dir / f"检验方法_{fixed_run_ts}_导入.xlsx",
                "检验方法#单据头(FBillHead)",
            )
            self.assertEqual(["1", "2"], [row["FBillHead(QM_InspectMethod)"] for row in method_rows])
            for row in method_rows:
                _assert_audit_fields(row, fixed_now)

            mod.gen_quality_std(
                pd.DataFrame([{"target": "Target A"}, {"target": "Target B"}]),
                {"QM_QualityStd": pd.DataFrame([{"FNumber": "ZLBZ000100", "FName#2052": "Old Target"}])},
            )
            std_rows = _row_maps(
                output_dir / f"质量标准_{fixed_run_ts}_导入.xlsx",
                "质量标准#单据头(FBillHead)",
            )
            self.assertEqual(["1", "2"], [row["FBillHead(QM_QualityStd)"] for row in std_rows])
            for row in std_rows:
                _assert_audit_fields(row, fixed_now)

            mod.gen_inspect_item(
                pd.DataFrame(
                    [
                        {
                            "test_item": "Item A",
                            "unit": "mg/kg",
                            "target": "Target A",
                            "method": "Method A",
                        },
                        {
                            "test_item": "Item B",
                            "unit": "g/L",
                            "target": "Target B",
                            "method": "Method B",
                        }
                    ]
                ),
                {"QM_InspectItem": pd.DataFrame([{"FNumber": "JYXM000100", "FName#2052": "Old Item"}])},
                {"mg/kg": "QA-001", "g/L": "QA-002"},
                {"Target A": "ZLBZ000101", "Target B": "ZLBZ000102"},
                {"Method A": "JYFF000101", "Method B": "JYFF000102"},
            )
            item_rows = _row_maps(
                output_dir / f"检验项目_{fixed_run_ts}_导入.xlsx",
                "检验项目#单据头(FBillHead)",
            )
            self.assertEqual(["1", "2"], [row["FBillHead(QM_InspectItem)"] for row in item_rows])
            for row in item_rows:
                _assert_audit_fields(row, fixed_now)

            mod.gen_qcscheme(
                pd.DataFrame(
                    [
                        {
                            "material_code": "301000001",
                            "material_name": "Material A",
                            "test_item": "Item A",
                            "target": "Target A",
                            "unit": "mg/kg",
                            "method": "Method A",
                        },
                        {
                            "material_code": "301000001",
                            "material_name": "Material A",
                            "test_item": "Item B",
                            "target": "Target B",
                            "unit": "g/L",
                            "method": "Method B",
                        },
                        {
                            "material_code": "301000002",
                            "material_name": "Material B",
                            "test_item": "Item A",
                            "target": "Target A",
                            "unit": "mg/kg",
                            "method": "Method A",
                        }
                    ]
                ),
                {"QM_QCScheme": pd.DataFrame([{"FNumber": "301000000"}])},
                {"Item A": "JYXM000101", "Item B": "JYXM000102"},
                {"Target A": "ZLBZ000101", "Target B": "ZLBZ000102"},
                {"Method A": "JYFF000101", "Method B": "JYFF000102"},
                {"mg/kg": "QA-001", "g/L": "QA-002"},
            )
            scheme_rows = _row_maps(
                output_dir / f"质检方案_{fixed_run_ts}_导入.xlsx",
                "质检方案#单据头(FBillHead)",
            )
            self.assertEqual(
                ["1", None, "2"],
                [row["FBillHead(QM_QCScheme)"] for row in scheme_rows],
            )
            self.assertIsNone(scheme_rows[0]["FEntity"])
            self.assertIsNone(scheme_rows[0]["FEntityScope"])
            self.assertIsNone(scheme_rows[0]["FMaterialId"])
            self.assertIsNone(scheme_rows[0]["FMaterialId#Name"])
            _assert_audit_fields(scheme_rows[0], fixed_now)
            _assert_audit_fields(scheme_rows[2], fixed_now)
        finally:
            mod.OUTPUT_DIR = original_output_dir
            mod.TODAY = original_today
            if original_now is None:
                delattr(mod, "NOW")
            else:
                mod.NOW = original_now
            mod.RUN_TS = original_run_ts
            for path in created_files:
                path.unlink(missing_ok=True)
