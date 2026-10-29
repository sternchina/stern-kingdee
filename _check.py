import openpyxl, os

target = os.path.join(os.getcwd(), 'workspace', 'BOM_API', 'output', 'K3_MaterialMaster.xlsx')
wb = openpyxl.load_workbook(target)
ws = wb.active
hdr = [c.value for c in ws[1]]

col = "full IBC quantity"
idx = hdr.index(col)
non_zero_vals = []
for row in ws.iter_rows(min_row=2, min_col=idx+1, max_col=idx+1):
    for c in row:
        if c.value is not None and str(c.value).strip() and str(c.value).strip() not in ('0', '0.0', ''):
            non_zero_vals.append(str(c.value))

with open('_result.txt', 'w', encoding='utf-8') as f:
    f.write(f"full IBC quantity non-zero count: {len(non_zero_vals)}\n")
    f.write(f"Sample: {non_zero_vals[:20]}\n")
